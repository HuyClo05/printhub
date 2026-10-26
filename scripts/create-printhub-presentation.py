from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = "Documentation/printcloud-one-page.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

NAVY = RGBColor(23, 32, 51)
MUTED = RGBColor(91, 105, 130)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(15, 159, 145)
LINE = RGBColor(218, 226, 239)
PALE = RGBColor(245, 248, 252)
ORANGE_PALE = RGBColor(255, 247, 237)
TEAL_PALE = RGBColor(239, 252, 249)
WHITE = RGBColor(255, 255, 255)

def box(x, y, w, h, fill=WHITE, line=LINE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line; s.line.width = Pt(.8)
    return s

def text(x, y, w, h, value, size=12, color=NAVY, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.04); tf.margin_right = Inches(.04)
    tf.margin_top = Inches(.02); tf.margin_bottom = Inches(.02)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; r.font.name = "Aptos"; r.font.size = Pt(size)
    r.font.bold = bold; r.font.color.rgb = color
    return tb

def step(x, y, w, label, fill, line):
    box(x, y, w, .72, fill, line)
    text(x + .04, y + .12, w - .08, .46, label, 8.5, NAVY, True, PP_ALIGN.CENTER)

slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
text(.62, .38, 3.5, .22, "PRINTCLOUD · CLOUD 3D PRINTING", 9, BLUE, True)
text(.62, .72, 7.8, .55, "PrintCloud", 30, NAVY, True)
text(.62, 1.31, 11.5, .42, "A cloud-hosted platform built around OrcaSlicer, with a dashboard to manage print jobs and printers.", 14, MUTED)
box(10.7, .62, 1.98, .45, RGBColor(233, 247, 245), RGBColor(185, 238, 230))
text(10.78, .74, 1.82, .18, "SLICE · SEND · MONITOR", 8, TEAL, True, PP_ALIGN.CENTER)

text(.62, 1.98, 7, .28, "THE 3D-PRINTING PIPELINE", 12, NAVY, True)
box(.62, 2.36, 6.02, 2.33, ORANGE_PALE, RGBColor(254, 215, 170))
box(6.72, 2.36, 6.0, 2.33, TEAL_PALE, RGBColor(185, 238, 230))
text(.88, 2.58, 5.5, .24, "Traditional workflow", 15, NAVY, True)
text(6.98, 2.58, 5.5, .24, "With PrintCloud", 15, NAVY, True)

xs = [.86, 2.02, 3.18, 4.34, 5.50]
labels = ["3D model", "Slice manually\nwith OrcaSlicer", "G-code\nfile", "Send by USB or\nlocal network", "3D printed\nmodel"]
for i, label in enumerate(labels):
    step(xs[i], 3.24, 1.0, label, WHITE, RGBColor(254, 215, 170))
    if i < 4: text(xs[i] + 1.01, 3.43, .14, .2, "→", 15, MUTED, True, PP_ALIGN.CENTER)
text(.88, 4.28, 5.52, .26, "Requires a slicer installed locally and access to the printer network.", 9.5, MUTED)

xs = [6.96, 8.12, 9.28, 10.44, 11.60]
labels = ["3D model", "PrintCloud\nauto-slice", "G-code\nfile", "Automatically sent\nto designated printer", "3D printed\nmodel"]
for i, label in enumerate(labels):
    step(xs[i], 3.24, 1.0, label, WHITE, RGBColor(185, 238, 230))
    if i < 4: text(xs[i] + 1.01, 3.43, .14, .2, "→", 15, MUTED, True, PP_ALIGN.CENTER)
text(6.98, 4.28, 5.48, .26, "Prepare and dispatch jobs from one cloud dashboard.", 9.5, MUTED)

text(.62, 5.05, 4, .28, "ADVANTAGES", 12, NAVY, True)
advantages = [
    ("No local slicer required", "Users do not need to download or install a 3D slicer to prepare models."),
    ("Print beyond the local network", "PrintCloud can send G-code to a designated printer outside the user’s printer network."),
]
for i, (heading, body) in enumerate(advantages):
    x = .62 + i * 6.1
    box(x, 5.48, 5.72, 1.08, PALE, LINE)
    text(x + .2, 5.68, 5.3, .22, "✓  " + heading, 12, TEAL if i else BLUE, True)
    text(x + .2, 6.04, 5.25, .31, body, 10, MUTED)

text(.62, 6.91, 12, .2, "PrintCloud connects model preparation and printer delivery into a simpler, more accessible workflow.", 10, NAVY, True, PP_ALIGN.CENTER)
prs.save(OUT)
print(OUT)
